"""Visualise the expert-evaluation ratings across eight DEMA subdimensions."""

from __future__ import annotations

from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from matplotlib.colors import BoundaryNorm, ListedColormap
from matplotlib.patches import Rectangle
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
INPUT_PATH = ROOT / "outputs" / "Expert evaluation of LLM-assisted DEMA assessments(1-1).xlsx"
OUTPUT_PATH = ROOT / "figures" / "expert_ratings_across_eight_subdimensions.png"

NAVY = "#123B52"
MUTED = "#607782"
GRID = "#FFFFFF"

# Microsoft Forms exported eight repeated blocks without their section titles.
# They are plotted in the standard DEMA subdimension order used by the form.
SUBDIMENSIONS = [
    "A.1  Roles",
    "A.2  Capacity",
    "B.1  Data supply",
    "B.2  Data infrastructure",
    "B.3  Data demand",
    "C.1  Governance, ethics & trust",
    "C.2  Policies, standards & resources",
    "C.3  Collaboration & inclusion",
]

QUESTIONS = [
    ("Q1", "Evidence relevance"),
    ("Q2", "Evidence sufficiency"),
    ("Q3", "Evidence specificity"),
    ("Q4", "Faithfulness / proportionality"),
    ("Q5", "Criterion alignment"),
    ("Q6", "Justification specificity"),
    ("Q7", "Practical usefulness"),
    ("Q8", "Diagnostic value"),
    ("Q9", "Phase defensibility"),
]

LIKERT = {
    "strongly disagree": 1,
    "disagree": 2,
    "neither agree nor disagree": 3,
    "agree": 4,
    "strongly agree": 5,
}


def normalise(value: object) -> str:
    return " ".join(str(value).replace("\n", " ").split()).strip().lower()


def read_ratings() -> np.ndarray:
    workbook = load_workbook(INPUT_PATH, data_only=True, read_only=True)
    sheet = workbook.active

    # Columns 1-7 are response metadata. Each subdimension then occupies ten
    # columns: nine Likert ratings followed by one categorical diagnostic item.
    values: list[list[int]] = []
    for block in range(8):
        first_column = 8 + block * 10
        row: list[int] = []
        for column in range(first_column, first_column + 9):
            raw = sheet.cell(row=2, column=column).value
            key = normalise(raw)
            if key not in LIKERT:
                raise ValueError(
                    f"Unexpected rating in row 2, column {column}: {raw!r}"
                )
            row.append(LIKERT[key])
        values.append(row)

    ratings = np.asarray(values, dtype=int)
    if ratings.shape != (8, 9):
        raise ValueError(f"Expected an 8 x 9 rating matrix, found {ratings.shape}")
    return ratings


def main() -> None:
    ratings = read_ratings()

    # Warm colours indicate disagreement; cool teal indicates agreement.
    colours = ["#B84A3A", "#E59A63", "#ECE8DC", "#73B7B0", "#176B87"]
    cmap = ListedColormap(colours)
    norm = BoundaryNorm([0.5, 1.5, 2.5, 3.5, 4.5, 5.5], cmap.N)

    plt.rcParams.update({"font.family": "DejaVu Sans"})
    fig, ax = plt.subplots(figsize=(11.2, 7.2), dpi=240)
    fig.patch.set_facecolor("white")

    fig.text(
        0.055,
        0.955,
        "Expert ratings across eight DEMA subdimensions",
        fontsize=19,
        fontweight="bold",
        color=NAVY,
        va="top",
    )
    fig.text(
        0.055,
        0.913,
        "One expert response · nine assessment questions · 1 = strongly disagree, 5 = strongly agree",
        fontsize=10.5,
        color=MUTED,
        va="top",
    )

    ax.imshow(ratings, cmap=cmap, norm=norm, aspect="auto")

    ax.set_xticks(range(9), labels=[code for code, _ in QUESTIONS])
    ax.set_yticks(range(8), labels=SUBDIMENSIONS)
    ax.tick_params(axis="x", top=True, bottom=False, labeltop=True, labelbottom=False,
                   length=0, pad=12, labelsize=10)
    ax.tick_params(axis="y", length=0, pad=11, labelsize=10)

    # White separators make every response readable at Word-page size.
    ax.set_xticks(np.arange(-0.5, 9, 1), minor=True)
    ax.set_yticks(np.arange(-0.5, 8, 1), minor=True)
    ax.grid(which="minor", color=GRID, linewidth=2.2)
    ax.tick_params(which="minor", bottom=False, left=False)
    for spine in ax.spines.values():
        spine.set_visible(False)

    abbreviations = {1: "SD", 2: "D", 3: "N", 4: "A", 5: "SA"}
    for row in range(ratings.shape[0]):
        for column in range(ratings.shape[1]):
            value = int(ratings[row, column])
            text_colour = "white" if value in (1, 5) else NAVY
            ax.text(
                column,
                row,
                f"{value}\n{abbreviations[value]}",
                ha="center",
                va="center",
                fontsize=9.2,
                fontweight="bold",
                color=text_colour,
                linespacing=0.9,
            )

    # Compact construct labels above the column labels.
    group_y = -1.20
    groups = [
        (-0.45, 2.45, "EVIDENCE PACKAGE"),
        (2.55, 7.45, "JUSTIFICATION"),
        (7.55, 8.45, "OVERALL"),
    ]
    for left, right, label in groups:
        ax.add_patch(
            Rectangle(
                (left, group_y - 0.16),
                right - left,
                0.32,
                facecolor="#EAF2F5",
                edgecolor="none",
                clip_on=False,
            )
        )
        ax.text(
            (left + right) / 2,
            group_y,
            label,
            ha="center",
            va="center",
            fontsize=7.5,
            fontweight="bold",
            color=LLM_BLUE,
            clip_on=False,
        )

    # Question key beneath the heatmap, split into two compact columns.
    key_left = [f"{code}  {label}" for code, label in QUESTIONS[:5]]
    key_right = [f"{code}  {label}" for code, label in QUESTIONS[5:]]
    fig.text(0.055, 0.205, "Question key", fontsize=10, fontweight="bold", color=NAVY)
    fig.text(0.055, 0.180, "\n".join(key_left), fontsize=8.4, color="#314852",
             va="top", linespacing=1.42)
    fig.text(0.50, 0.180, "\n".join(key_right), fontsize=8.4, color="#314852",
             va="top", linespacing=1.42)

    legend_labels = [
        "1  Strongly disagree",
        "2  Disagree",
        "3  Neither",
        "4  Agree",
        "5  Strongly agree",
    ]
    legend_handles = [
        Rectangle((0, 0), 1, 1, facecolor=colour, edgecolor="none")
        for colour in colours
    ]
    fig.legend(
        legend_handles,
        legend_labels,
        loc="lower center",
        bbox_to_anchor=(0.5, 0.055),
        ncol=5,
        frameon=False,
        fontsize=8.2,
        handlelength=1.2,
        columnspacing=1.5,
    )
    fig.text(
        0.055,
        0.018,
        "Note. The Excel export contains one completed response. Rows follow the eight DEMA subdimensions in framework order; "
        "the export itself does not retain Microsoft Forms section titles.",
        fontsize=7.5,
        color=MUTED,
    )

    plt.subplots_adjust(left=0.275, right=0.965, top=0.77, bottom=0.29)
    fig.savefig(OUTPUT_PATH, facecolor="white", bbox_inches="tight", pad_inches=0.12)
    plt.close(fig)
    print(OUTPUT_PATH)


LLM_BLUE = "#176B87"


if __name__ == "__main__":
    main()
